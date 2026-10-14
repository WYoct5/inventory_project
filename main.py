import csv
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.styles import PatternFill
from datetime import datetime


stock_data = {}

with open("data/stock.csv",encoding="utf-8") as f:
    reader = csv.DictReader(f)

    for row in reader:
        stock = int(row["stock"])
        stock_data[row["product"]] = stock

    print(stock_data)



shipment_totals = {}
shipment_dates = []


with open("data/shipments.csv",encoding="utf-8") as f:
    reader = csv.DictReader(f)

    for row in reader:
        product = row["product"]
        quantity = int(row["quantity"])
        date = datetime.strptime(row["date"],"%Y-%m-%d")

        shipment_dates.append(date)

        if product in shipment_totals:
            shipment_totals[product] += quantity
        else:
            shipment_totals[product] = quantity

if shipment_dates:
    first_date = min(shipment_dates)
    last_date = max(shipment_dates)
    days = (last_date - first_date).days + 1
else:
    days = 0

print("集計期間:",days,"日")


lead_times = {
    "JN-CA01":10,
    "JN-CB02":5,
    "AB-CD03":7,
    "XY-ZA04":3
}


result = []

for product,stock in stock_data.items():
    
    if product in shipment_totals:
        shipment = shipment_totals[product]
    else:
        shipment = 0

    if days > 0:
        average_daily_shipment = shipment / days
    else:
        average_daily_shipment = 0

    if average_daily_shipment > 0:
        days_remaining = stock / average_daily_shipment
    else:
        days_remaining = None
        risk = "出荷なし"

    lead_time = lead_times.get(product)

    if days_remaining is None:
        risk = "出荷なし"
    elif days_remaining <= lead_time:
        risk = "危険"
    elif days_remaining <= lead_time + 7:
        risk = "注意"
    else:
        risk = "問題なし"


    

    result.append({
        "product": product,
        "stock": stock,
        "shipment": shipment,
        "average_daily_shipment": average_daily_shipment,
        "days_remaining": days_remaining,
        "risk": risk
    })

wb = Workbook()
ws = wb.active

for row in result:

    days_remaining = row["days_remaining"]

    if days_remaining is None:
        days_remaining = "-"
    else:
        days_remaining = round(days_remaining,1)


    print(
        row["product"],
        row["stock"],
        row["shipment"],
        round(row["average_daily_shipment"],1),
        days_remaining,
        row["risk"]
    )


    ws.append([
        "商品コード",
        "現在庫",
        "出荷合計",
        "1日平均出荷数",
        "残り日数",
        "リスク",
    ])

    for cell in ws[1]:
        cell.font = Font(bold = True)

    for row in result:
        ws.append([
            row["product"],
            row["stock"],
            row["shipment"],
            round(row["average_daily_shipment"],1),
            row["days_remaining"],
            row["risk"]
        ])

for row_num in range(2,ws.max_row + 1):
    risk = ws.cell(row = row_num,column = 6).value

    if risk == "危険":
        color = "FF9999"
    elif risk == "注意":
        color = "FFFF99"
    elif risk == "問題なし":
        color = "99FF99"
    elif risk == "リードタイム未登録":
        color = "CCCCCC"
    else:
        color = None

    if color :
        for cell in ws[row_num]:
            cell.fill = PatternFill(
                fill_type = "solid",
                fgColor = color
            )
for column in ws.columns:
    max_length = 0

    for cell in column:
        if cell.value is not None:
            max_length = max(max_length,len(str(cell.value)))

        ws.column_dimensions[column[0].column_letter].width = max_length + 5


wb.save("inventory_result.xlsx")
print("Excelファイルを保存しました")